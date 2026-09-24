from abc import ABC, abstractmethod
from pathlib import Path
from .domain import CanonicalDocument , Segment, SourceType, SourceFile
import pymupdf
from pptx import Presentation
import re

class SourceAdapter(ABC):

    @abstractmethod
    def parse(
        self,
        source: SourceFile,
    ) -> CanonicalDocument:
        pass

class TXTAdapter(SourceAdapter):

    def parse(
        self,
        source: SourceFile,
    ) -> CanonicalDocument:

        text = source.path.read_text(
            encoding="utf-8"
        ).strip()

        segments = []

        if text:
            segments.append(
                Segment(text=text)
            )

        return CanonicalDocument(
            filename=source.original_filename,
            source_type=SourceType.TXT,
            segments=segments,
        )
    
class PDFAdapter(SourceAdapter):

    def parse(
        self,
        source: SourceFile,
    ) -> CanonicalDocument:

        try:
            document = pymupdf.open(
                source.path
            )

            segments = []

            for page_number, page in enumerate(
                document,
                start=1,
            ):
                text = page.get_text(
                    "text"
                ).strip()

                if text:
                    segments.append(
                        Segment(
                            text=text,
                            page_number=page_number,
                        )
                    )

            document.close()

            return CanonicalDocument(
                filename=source.original_filename,
                source_type=SourceType.PDF,
                segments=segments,
            )

        except Exception as exc:
            raise DocumentParsingError(
                f"Failed to parse {source.original_filename}"
            ) from exc

class DocumentParsingError(Exception):
    pass


class UnsupportedSourceTypeError(Exception):
    pass

class AdapterFactory:

    @staticmethod
    def get_adapter(
        source: SourceFile,
    ) -> SourceAdapter:

        extension = (
            source.path.suffix.lower()
        )

        if extension == ".txt":
            return TXTAdapter()

        if extension == ".pdf":
            return PDFAdapter()

        if extension == ".pptx":
            return PPTXAdapter()
        
        if extension == ".vtt":
            return VTTAdapter()

        raise UnsupportedSourceTypeError(
            f"Unsupported source type: {extension}"
        )

class PPTXAdapter(SourceAdapter):

    def parse(
        self,
        source: SourceFile,
    ) -> CanonicalDocument:

        presentation = Presentation(
            source.path
        )

        segments = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):

            texts = []

            for shape in slide.shapes:

                text = getattr(
                    shape,
                    "text",
                    "",
                ).strip()

                if text:
                    texts.append(text)

            combined_text = "\n".join(
                texts
            )

            if not combined_text:
                continue

            segments.append(
                Segment(
                    text=combined_text,
                    slide_number=slide_number,
                )
            )

        return CanonicalDocument(
            filename=source.original_filename,
            source_type=SourceType.PPTX,
            segments=segments,
        )
def timestamp_to_seconds(
    value: str,
) -> float:

    hours, minutes, seconds = (
        value.split(":")
    )

    return (
        int(hours) * 3600
        + int(minutes) * 60
        + float(seconds)
    )

class VTTAdapter(SourceAdapter):

    TIMESTAMP_PATTERN = re.compile(
        r"(\d{2}:\d{2}:\d{2}\.\d{3})"
        r"\s+-->\s+"
        r"(\d{2}:\d{2}:\d{2}\.\d{3})"
    )

    def parse(
        self,
        source: SourceFile,
    ) -> CanonicalDocument:

        content = source.path.read_text(
            encoding="utf-8"
        )

        lines = content.splitlines()

        segments = []

        current_start = None
        current_end = None
        current_text = []

        for raw_line in lines:

            line = raw_line.strip()

            match = (
                self.TIMESTAMP_PATTERN.match(
                    line
                )
            )

            if match:

                if current_text:

                    segments.append(
                        Segment(
                            text=" ".join(
                                current_text
                            ),
                            timestamp_start=(
                                current_start
                            ),
                            timestamp_end=(
                                current_end
                            ),
                        )
                    )

                    current_text = []

                current_start = (
                    timestamp_to_seconds(
                        match.group(1)
                    )
                )

                current_end = (
                    timestamp_to_seconds(
                        match.group(2)
                    )
                )

                continue

            if (
                not line
                or line == "WEBVTT"
            ):
                continue

            current_text.append(line)

        if current_text:

            segments.append(
                Segment(
                    text=" ".join(
                        current_text
                    ),
                    timestamp_start=(
                        current_start
                    ),
                    timestamp_end=(
                        current_end
                    ),
                )
            )

        return CanonicalDocument(
            filename=source.original_filename,
            source_type=SourceType.VTT,
            segments=segments,
        )