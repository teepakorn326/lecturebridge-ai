from app.modules.ingestion.adapters import (
    TXTAdapter,VTTAdapter,PDFAdapter,DocumentParsingError,PPTXAdapter,AdapterFactory,UnsupportedSourceTypeError
)
from app.modules.ingestion.domain import (
    SourceFile,
    SourceType,
)

import pymupdf
import pytest

from pptx import Presentation
from pptx.util import Inches


def test_txt_adapter(tmp_path):

    path = tmp_path / "lecture.txt"

    path.write_text(
        "Attention mechanism.",
        encoding="utf-8",
    )

    source = SourceFile(
        path=path,
        original_filename="lecture.txt",
    )

    adapter = TXTAdapter()

    document = adapter.parse(
        source
    )

    assert document.filename == (
        "lecture.txt"
    )

    assert (
        document.source_type
        == SourceType.TXT
    )

    assert len(
        document.segments
    ) == 1

    assert (
        document.segments[0].text
        == "Attention mechanism."
    )

def test_vtt_adapter(tmp_path):

    file_path = tmp_path / "lecture.vtt"

    file_path.write_text(
        """WEBVTT

00:00:01.000 --> 00:00:04.000
Welcome to the lecture.

00:00:05.000 --> 00:00:09.000
Today we study attention.
""",
        encoding="utf-8",
    )

    source = SourceFile(
        path=file_path,
        original_filename="lecture.vtt",
    )

    adapter = VTTAdapter()

    document = adapter.parse(source)

    assert document.source_type == SourceType.VTT

    assert len(document.segments) == 2

    first = document.segments[0]

    assert first.text == "Welcome to the lecture."
    assert first.timestamp_start == 1.0
    assert first.timestamp_end == 4.0

def test_pdf_adapter(tmp_path):
    # สร้าง PDF 2 หน้าแบบ programmatic
    pdf_path = tmp_path / "lecture.pdf"

    doc = pymupdf.open()                      # PDF เปล่า
    page1 = doc.new_page()
    page1.insert_text((72, 72), "Gradient descent basics")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Backpropagation explained")
    doc.save(pdf_path)
    doc.close()

    source = SourceFile(path=pdf_path, original_filename="lecture.pdf")

    document = PDFAdapter().parse(source)

    assert document.source_type == SourceType.PDF
    assert len(document.segments) == 2
    assert document.segments[0].page_number == 1        # ← citation ต้องรอด!
    assert "Gradient descent" in document.segments[0].text
    assert document.segments[1].page_number == 2

def test_pdf_adapter_corrupted_file(tmp_path):
    bad_path = tmp_path / "broken.pdf"
    bad_path.write_bytes(b"this is not a real pdf")

    source = SourceFile(path=bad_path, original_filename="broken.pdf")

    with pytest.raises(DocumentParsingError):
        PDFAdapter().parse(source)

def test_pptx_adapter(tmp_path):
    pptx_path = tmp_path / "lecture.pptx"

    prs = Presentation()
    blank = prs.slide_layouts[6]              # layout เปล่า

    slide1 = prs.slides.add_slide(blank)
    box = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Attention mechanism"

    slide2 = prs.slides.add_slide(blank)      # สไลด์เปล่า — ต้องถูกข้าม

    slide3 = prs.slides.add_slide(blank)
    box3 = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box3.text_frame.text = "Transformer architecture"

    prs.save(pptx_path)

    source = SourceFile(path=pptx_path, original_filename="lecture.pptx")
    document = PPTXAdapter().parse(source)

    assert document.source_type == SourceType.PPTX
    assert len(document.segments) == 2                 # สไลด์เปล่าถูกข้าม
    assert document.segments[0].slide_number == 1
    assert document.segments[1].slide_number == 3      # ← ไม่ใช่ 2!

def test_factory_selects_adapter_by_extension(tmp_path):
    txt = SourceFile(path=tmp_path / "a.txt", original_filename="a.txt")
    vtt = SourceFile(path=tmp_path / "a.vtt", original_filename="a.vtt")

    assert isinstance(AdapterFactory.get_adapter(txt), TXTAdapter)
    assert isinstance(AdapterFactory.get_adapter(vtt), VTTAdapter)


def test_factory_rejects_unknown_extension(tmp_path):
    source = SourceFile(path=tmp_path / "a.docx", original_filename="a.docx")

    with pytest.raises(UnsupportedSourceTypeError):
        AdapterFactory.get_adapter(source)